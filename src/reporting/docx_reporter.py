"""
Professional Word Document (DOCX) Report Generator

Generates editable JORC-compliant QAQC analysis reports with:
- Cover page with project branding
- Executive summary with pass/fail indicators
- Detailed analysis sections with editable tables
- JORC Table 1 helper text with edit prompts
"""

import os
from datetime import datetime
from typing import Dict, List, Optional, Any

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class QAQCDOCXReporter:
    """
    Professional Word document report generator for QAQC analysis results.
    
    Creates editable documents that can be customized by the user
    for JORC compliance submissions and Competent Person review.
    """
    
    # Color scheme (RGB values)
    COLORS = {
        'primary': RGBColor(245, 158, 11),      # Amber/Gold
        'primary_dark': RGBColor(217, 119, 6),
        'success': RGBColor(16, 185, 129),       # Green
        'error': RGBColor(239, 68, 68),          # Red
        'text_dark': RGBColor(30, 41, 59),       # Slate 800
        'text_light': RGBColor(100, 116, 139),   # Slate 500
        'border': RGBColor(226, 232, 240),       # Slate 200
    }
    
    def __init__(self, config: Dict = None) -> None:
        """
        Initialize DOCX reporter with configuration.
        
        Args:
            config: Dictionary with report settings
        """
        self.config = config or {}
        self.include_plots = self.config.get('include_plots', True)
        self.document = None
    
    def _create_document(self) -> Document:
        """Create a new document with custom styles."""
        doc = Document()
        
        # Set default font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        font.color.rgb = self.COLORS['text_dark']
        
        # Create custom styles
        self._add_custom_styles(doc)
        
        return doc
    
    def _add_custom_styles(self, doc: Document) -> None:
        """Add custom paragraph and table styles."""
        styles = doc.styles
        
        # Title style
        if 'QAQC Title' not in [s.name for s in styles]:
            title_style = styles.add_style('QAQC Title', WD_STYLE_TYPE.PARAGRAPH)
            title_style.font.size = Pt(28)
            title_style.font.bold = True
            title_style.font.color.rgb = self.COLORS['text_dark']
            title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_style.paragraph_format.space_after = Pt(6)
        
        # Subtitle style
        if 'QAQC Subtitle' not in [s.name for s in styles]:
            subtitle_style = styles.add_style('QAQC Subtitle', WD_STYLE_TYPE.PARAGRAPH)
            subtitle_style.font.size = Pt(14)
            subtitle_style.font.color.rgb = self.COLORS['text_light']
            subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle_style.paragraph_format.space_after = Pt(24)
        
        # Section heading style
        if 'QAQC Heading' not in [s.name for s in styles]:
            heading_style = styles.add_style('QAQC Heading', WD_STYLE_TYPE.PARAGRAPH)
            heading_style.font.size = Pt(16)
            heading_style.font.bold = True
            heading_style.font.color.rgb = self.COLORS['primary_dark']
            heading_style.paragraph_format.space_before = Pt(18)
            heading_style.paragraph_format.space_after = Pt(10)
        
        # Subsection heading style
        if 'QAQC Subheading' not in [s.name for s in styles]:
            subheading_style = styles.add_style('QAQC Subheading', WD_STYLE_TYPE.PARAGRAPH)
            subheading_style.font.size = Pt(12)
            subheading_style.font.bold = True
            subheading_style.font.color.rgb = self.COLORS['text_dark']
            subheading_style.paragraph_format.space_before = Pt(12)
            subheading_style.paragraph_format.space_after = Pt(6)
        
        # Status Pass style
        if 'Status Pass' not in [s.name for s in styles]:
            pass_style = styles.add_style('Status Pass', WD_STYLE_TYPE.PARAGRAPH)
            pass_style.font.size = Pt(14)
            pass_style.font.bold = True
            pass_style.font.color.rgb = self.COLORS['success']
            pass_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Status Fail style
        if 'Status Fail' not in [s.name for s in styles]:
            fail_style = styles.add_style('Status Fail', WD_STYLE_TYPE.PARAGRAPH)
            fail_style.font.size = Pt(14)
            fail_style.font.bold = True
            fail_style.font.color.rgb = self.COLORS['error']
            fail_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Edit prompt style
        if 'Edit Prompt' not in [s.name for s in styles]:
            edit_style = styles.add_style('Edit Prompt', WD_STYLE_TYPE.PARAGRAPH)
            edit_style.font.size = Pt(10)
            edit_style.font.italic = True
            edit_style.font.color.rgb = self.COLORS['primary']
    
    def _add_horizontal_line(self, doc: Document) -> None:
        """Add a horizontal line to the document."""
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        
        # Create a bottom border
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'F59E0B')  # Amber color
        pBdr.append(bottom)
        pPr.append(pBdr)
    
    def _create_table(self, doc: Document, headers: List[str], rows: List[List[str]]) -> None:
        """Create a styled table with headers and data rows."""
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header row
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            cell = header_cells[i]
            cell.text = header
            # Style header
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.runs[0]
            run.font.bold = True
            run.font.size = Pt(10)
            # Set header background color
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F59E0B')  # Amber
            cell._tc.get_or_add_tcPr().append(shading)
            run.font.color.rgb = RGBColor(255, 255, 255)
        
        # Data rows
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, cell_text in enumerate(row_data):
                cell = row_cells[col_idx]
                cell.text = str(cell_text)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                run.font.size = Pt(10)
                
                # Alternate row shading
                if row_idx % 2 == 1:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'F8FAFC')  # Light gray
                    cell._tc.get_or_add_tcPr().append(shading)
        
        doc.add_paragraph()  # Add spacing after table
    
    def _create_metrics_table(self, doc: Document, metrics: List[Dict]) -> None:
        """Create a two-column metrics table."""
        table = doc.add_table(rows=len(metrics), cols=2)
        table.style = 'Table Grid'
        
        for row_idx, metric in enumerate(metrics):
            cells = table.rows[row_idx].cells
            
            # Label cell
            cells[0].text = metric['label']
            cells[0].paragraphs[0].runs[0].font.bold = True
            cells[0].paragraphs[0].runs[0].font.size = Pt(10)
            cells[0].paragraphs[0].runs[0].font.color.rgb = self.COLORS['text_light']
            
            # Value cell
            cells[1].text = str(metric['value'])
            run = cells[1].paragraphs[0].runs[0]
            run.font.size = Pt(11)
            run.font.bold = True
            
            # Color based on status
            if metric.get('status') == 'pass':
                run.font.color.rgb = self.COLORS['success']
            elif metric.get('status') == 'fail':
                run.font.color.rgb = self.COLORS['error']
            else:
                run.font.color.rgb = self.COLORS['text_dark']
        
        doc.add_paragraph()
    
    def build_cover_page(self, doc: Document, project_info: Dict) -> None:
        """Build the cover page."""
        # Add spacing at top
        for _ in range(4):
            doc.add_paragraph()
        
        # Title
        title = doc.add_paragraph('QAQC ANALYSIS REPORT', style='QAQC Title')
        
        # Subtitle with project info
        project_name = project_info.get('name', 'Untitled Project')
        deposit = project_info.get('deposit', '')
        subtitle_text = project_name
        if deposit:
            subtitle_text += f" | {deposit}"
        doc.add_paragraph(subtitle_text, style='QAQC Subtitle')
        
        # Horizontal line
        self._add_horizontal_line(doc)
        
        # Project details table
        doc.add_paragraph()
        details = [
            ('Commodity:', project_info.get('commodity', 'N/A')),
            ('Analysis Date:', datetime.now().strftime('%B %d, %Y')),
            ('Report Generated:', datetime.now().strftime('%Y-%m-%d %H:%M')),
        ]
        
        if project_info.get('laboratory'):
            details.append(('Laboratory:', project_info['laboratory']))
        if project_info.get('competent_person'):
            details.append(('Competent Person:', project_info['competent_person']))
        
        table = doc.add_table(rows=len(details), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for row_idx, (label, value) in enumerate(details):
            cells = table.rows[row_idx].cells
            cells[0].text = label
            cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            cells[0].paragraphs[0].runs[0].font.bold = True
            cells[0].width = Inches(1.5)
            
            cells[1].text = value
            cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            cells[1].width = Inches(3.0)
        
        # Page break
        doc.add_page_break()
    
    def build_executive_summary(self, doc: Document, analysis_results: Dict) -> None:
        """Build the executive summary section."""
        doc.add_paragraph('EXECUTIVE SUMMARY', style='QAQC Heading')
        self._add_horizontal_line(doc)
        
        # Get overall status
        standards = analysis_results.get('standards', {})
        blanks = analysis_results.get('blanks', {})
        duplicates = analysis_results.get('duplicates', {})
        
        overall_pass = all([
            standards.get('overall_acceptable', True),
            blanks.get('overall_acceptable', True),
            duplicates.get('overall_acceptable', True)
        ])
        
        # Overall status
        doc.add_paragraph('Overall QAQC Status', style='QAQC Subheading')
        status_style = 'Status Pass' if overall_pass else 'Status Fail'
        status_text = '✓ PASS' if overall_pass else '✗ FAIL'
        doc.add_paragraph(status_text, style=status_style)
        
        doc.add_paragraph()
        
        # Summary table
        headers = ['Analysis Component', 'Status', 'Key Metric', 'Value']
        rows = [
            [
                'Standards (CRMs)',
                'PASS' if standards.get('overall_acceptable', True) else 'FAIL',
                'RSD',
                f"{standards.get('precision', {}).get('rsd', 0):.1f}%"
            ],
            [
                'Blanks',
                'PASS' if blanks.get('overall_acceptable', True) else 'FAIL',
                'Contamination Rate',
                f"{blanks.get('contamination', {}).get('contamination_rate', 0) * 100:.1f}%"
            ],
            [
                'Duplicates',
                'PASS' if duplicates.get('overall_acceptable', True) else 'FAIL',
                'Mean RPD',
                f"{duplicates.get('precision', {}).get('mean_rpd', 0):.1f}%"
            ],
        ]
        self._create_table(doc, headers, rows)
        
        # Recommendations
        doc.add_paragraph('Recommendations', style='QAQC Subheading')
        recommendations = self._generate_recommendations(analysis_results)
        for rec in recommendations:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(rec)
        
        doc.add_page_break()
    
    def _generate_recommendations(self, analysis_results: Dict) -> List[str]:
        """Generate recommendations based on analysis results."""
        recommendations = []
        
        standards = analysis_results.get('standards', {})
        blanks = analysis_results.get('blanks', {})
        duplicates = analysis_results.get('duplicates', {})
        
        if standards.get('bias', {}).get('bias_detected', False):
            recommendations.append("Investigate and correct systematic bias in standards analysis")
        
        if not blanks.get('contamination', {}).get('acceptable', True):
            recommendations.append("Investigate contamination sources in blank samples")
        
        if blanks.get('carryover', {}).get('carryover_detected', False):
            recommendations.append("Implement additional rinsing between samples to reduce carry-over")
        
        if not duplicates.get('precision', {}).get('acceptable', True):
            recommendations.append("Improve duplicate precision through better homogenization")
        
        if duplicates.get('nugget_ratio', 0) > 0.3:
            recommendations.append("Consider spatial sampling strategy improvements due to high nugget ratio")
        
        if not recommendations:
            recommendations.append("All QAQC parameters are within acceptable limits. Continue current protocols.")
        
        return recommendations
    
    def build_standards_section(self, doc: Document, standards_results: Dict, plot_path: str = None) -> None:
        """Build the standards analysis section."""
        doc.add_paragraph('STANDARDS ANALYSIS', style='QAQC Heading')
        self._add_horizontal_line(doc)
        
        # Description
        p = doc.add_paragraph()
        p.add_run(
            "Certified Reference Materials (CRMs) are used to assess analytical accuracy and detect bias. "
            "The following results show the performance of standards against certified values."
        )
        
        # Status
        overall_pass = standards_results.get('overall_acceptable', True)
        status_style = 'Status Pass' if overall_pass else 'Status Fail'
        status_text = '✓ PASS' if overall_pass else '✗ FAIL'
        doc.add_paragraph(status_text, style=status_style)
        
        doc.add_paragraph()
        
        # Metrics
        bias = standards_results.get('bias', {})
        recovery = standards_results.get('recovery', {})
        precision = standards_results.get('precision', {})
        
        metrics = [
            {'label': 'Bias Detected', 'value': 'Yes' if bias.get('bias_detected', False) else 'No',
             'status': 'fail' if bias.get('bias_detected', False) else 'pass'},
            {'label': 'Max Z-Score', 'value': f"{bias.get('max_z_score', 0):.2f}"},
            {'label': 'Mean Recovery', 'value': f"{recovery.get('mean_recovery', 100):.1f}%",
             'status': 'pass' if recovery.get('acceptable', True) else 'fail'},
            {'label': 'Precision (RSD)', 'value': f"{precision.get('rsd', 0):.2f}%",
             'status': 'pass' if precision.get('acceptable', True) else 'fail'},
        ]
        self._create_metrics_table(doc, metrics)
        
        # Plot placeholder
        if plot_path and os.path.exists(plot_path):
            doc.add_paragraph('Control Chart', style='QAQC Subheading')
            doc.add_picture(plot_path, width=Inches(5.5))
        else:
            doc.add_paragraph('[Insert Control Chart Here]', style='Edit Prompt')
        
        doc.add_page_break()
    
    def build_blanks_section(self, doc: Document, blanks_results: Dict, plot_path: str = None) -> None:
        """Build the blanks analysis section."""
        doc.add_paragraph('BLANKS ANALYSIS', style='QAQC Heading')
        self._add_horizontal_line(doc)
        
        p = doc.add_paragraph()
        p.add_run(
            "Blank samples are analyzed to detect contamination and assess background levels. "
            "Contamination is flagged when blank values exceed the detection limit multiplied by a threshold factor."
        )
        
        # Status
        overall_pass = blanks_results.get('overall_acceptable', True)
        status_style = 'Status Pass' if overall_pass else 'Status Fail'
        status_text = '✓ PASS' if overall_pass else '✗ FAIL'
        doc.add_paragraph(status_text, style=status_style)
        
        doc.add_paragraph()
        
        contamination = blanks_results.get('contamination', {})
        carryover = blanks_results.get('carryover', {})
        background = blanks_results.get('background', {})
        
        metrics = [
            {'label': 'Contamination Rate', 'value': f"{contamination.get('contamination_rate', 0) * 100:.1f}%",
             'status': 'pass' if contamination.get('acceptable', True) else 'fail'},
            {'label': 'Method Detection Limit', 'value': f"{blanks_results.get('mdl', 0):.4f}"},
            {'label': 'Carry-over Detected', 'value': 'Yes' if carryover.get('carryover_detected', False) else 'No',
             'status': 'fail' if carryover.get('carryover_detected', False) else 'pass'},
            {'label': 'Background Mean', 'value': f"{background.get('mean', 0):.4f}"},
        ]
        self._create_metrics_table(doc, metrics)
        
        if plot_path and os.path.exists(plot_path):
            doc.add_paragraph('Blanks Distribution', style='QAQC Subheading')
            doc.add_picture(plot_path, width=Inches(5.5))
        else:
            doc.add_paragraph('[Insert Blanks Histogram Here]', style='Edit Prompt')
        
        doc.add_page_break()
    
    def build_duplicates_section(self, doc: Document, duplicates_results: Dict, plot_path: str = None) -> None:
        """Build the duplicates analysis section."""
        doc.add_paragraph('DUPLICATES ANALYSIS', style='QAQC Heading')
        self._add_horizontal_line(doc)
        
        p = doc.add_paragraph()
        p.add_run(
            "Duplicate samples assess analytical and sampling precision. Results are evaluated using "
            "Relative Percent Difference (RPD) and Half Absolute Relative Difference (HARD) metrics."
        )
        
        # Status
        overall_pass = duplicates_results.get('overall_acceptable', True)
        status_style = 'Status Pass' if overall_pass else 'Status Fail'
        status_text = '✓ PASS' if overall_pass else '✗ FAIL'
        doc.add_paragraph(status_text, style=status_style)
        
        doc.add_paragraph()
        
        precision = duplicates_results.get('precision', {})
        systematic = duplicates_results.get('systematic_errors', {})
        
        metrics = [
            {'label': 'Mean RPD', 'value': f"{precision.get('mean_rpd', 0):.1f}%",
             'status': 'pass' if precision.get('acceptable', True) else 'fail'},
            {'label': 'Max RPD', 'value': f"{precision.get('max_rpd', 0):.1f}%"},
            {'label': 'Systematic Error', 'value': 'Yes' if systematic.get('systematic_error', False) else 'No',
             'status': 'fail' if systematic.get('systematic_error', False) else 'pass'},
            {'label': 'Nugget Ratio', 'value': f"{duplicates_results.get('nugget_ratio', 0):.3f}"},
        ]
        self._create_metrics_table(doc, metrics)
        
        if plot_path and os.path.exists(plot_path):
            doc.add_paragraph('Duplicates Scatter Plot', style='QAQC Subheading')
            doc.add_picture(plot_path, width=Inches(5.5))
        else:
            doc.add_paragraph('[Insert Duplicates Scatter Plot Here]', style='Edit Prompt')
        
        doc.add_page_break()
    
    def build_jorc_section(self, doc: Document, analysis_results: Dict) -> None:
        """Build the editable JORC Table 1 helper section."""
        doc.add_paragraph('JORC TABLE 1 HELPER', style='QAQC Heading')
        self._add_horizontal_line(doc)
        
        # Edit notice
        notice = doc.add_paragraph(style='Edit Prompt')
        notice.add_run(
            "The following text is provided to assist the Competent Person in completing JORC Table 1, "
            "Section 1 (Sampling Techniques and Data). Please review and modify as appropriate."
        )
        
        doc.add_paragraph()
        
        # Criteria heading
        doc.add_paragraph('Criteria: Quality of assay data and laboratory tests', style='QAQC Subheading')
        
        # JORC requirement quote
        quote = doc.add_paragraph()
        quote_run = quote.add_run(
            '"Nature of quality control procedures adopted (e.g. standards, blanks, duplicates, '
            'external laboratory checks) and whether acceptable levels of accuracy (i.e. lack of bias) '
            'and precision have been established."'
        )
        quote_run.font.italic = True
        quote_run.font.size = Pt(10)
        quote_run.font.color.rgb = self.COLORS['text_light']
        
        doc.add_paragraph()
        
        # Edit prompt
        doc.add_paragraph('[EDIT THE TEXT BELOW TO MATCH YOUR PROJECT SPECIFICS]', style='Edit Prompt')
        
        doc.add_paragraph()
        
        # Generate suggested text
        standards = analysis_results.get('standards', {})
        blanks = analysis_results.get('blanks', {})
        duplicates = analysis_results.get('duplicates', {})
        
        # Suggested text paragraph
        doc.add_paragraph('SUGGESTED TEXT:', style='QAQC Subheading')
        
        suggested = doc.add_paragraph()
        suggested.add_run(
            "A comprehensive QAQC program was implemented including certified reference materials (standards), "
            "blank samples, and field duplicates."
        )
        
        doc.add_paragraph()
        
        # Standards
        std_para = doc.add_paragraph()
        std_para.add_run("Standards: ").bold = True
        standards_pass = standards.get('overall_acceptable', True)
        std_para.add_run(
            f"Analysis of CRMs indicates that analytical accuracy is "
            f"{'acceptable with no significant bias detected' if standards_pass else 'outside acceptable limits, with bias detected'}. "
            f"Mean recovery is {standards.get('recovery', {}).get('mean_recovery', 100):.1f}%."
        )
        
        # Blanks
        blk_para = doc.add_paragraph()
        blk_para.add_run("Blanks: ").bold = True
        blanks_pass = blanks.get('overall_acceptable', True)
        blk_para.add_run(
            f"Blank sample analysis indicates that contamination is "
            f"{'minimal and within acceptable limits' if blanks_pass else 'present and requires investigation'}. "
            f"Contamination rate is {blanks.get('contamination', {}).get('contamination_rate', 0) * 100:.1f}%."
        )
        
        # Duplicates
        dup_para = doc.add_paragraph()
        dup_para.add_run("Duplicates: ").bold = True
        duplicates_pass = duplicates.get('overall_acceptable', True)
        dup_para.add_run(
            f"Duplicate analysis indicates that sampling and analytical precision is "
            f"{'acceptable' if duplicates_pass else 'outside acceptable limits'}. "
            f"Mean RPD is {duplicates.get('precision', {}).get('mean_rpd', 0):.1f}%."
        )
        
        doc.add_paragraph()
        
        # Conclusion
        overall_pass = all([standards_pass, blanks_pass, duplicates_pass])
        conclusion = doc.add_paragraph()
        if overall_pass:
            conclusion.add_run("Overall, acceptable levels of accuracy and precision have been established.")
        else:
            conclusion.add_run("Further investigation is recommended to address identified issues.")
        
        doc.add_paragraph()
        
        # Additional criteria
        doc.add_paragraph('Criteria: Verification of sampling and assaying', style='QAQC Subheading')
        
        quote2 = doc.add_paragraph()
        quote2_run = quote2.add_run(
            '"The verification of significant intersections by either independent or alternative personnel."'
        )
        quote2_run.font.italic = True
        quote2_run.font.size = Pt(10)
        quote2_run.font.color.rgb = self.COLORS['text_light']
        
        doc.add_paragraph()
        doc.add_paragraph('[EDIT THE TEXT BELOW TO MATCH YOUR PROJECT SPECIFICS]', style='Edit Prompt')
        doc.add_paragraph()
        
        doc.add_paragraph('SUGGESTED TEXT:', style='QAQC Subheading')
        verify = doc.add_paragraph()
        verify.add_run(
            "Significant intersections have been verified by internal QAQC procedures. "
            "[MODIFY: Describe any independent verification performed, or state if not applicable.]"
        )
    
    def generate_docx_report(
        self,
        analysis_results: Dict,
        project_info: Dict = None,
        plots: Dict = None,
        filename: str = None
    ) -> str:
        """
        Generate complete Word document report.
        
        Args:
            analysis_results: Dictionary with all analysis results
            project_info: Dictionary with project metadata
            plots: Dictionary with plot file paths
            filename: Output filename (optional)
            
        Returns:
            Path to generated DOCX file
        """
        # Generate filename if not provided
        if filename is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"qaqc_report_{timestamp}.docx"
        
        # Ensure output directory exists
        output_dir = os.path.dirname(filename)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        
        # Create document
        doc = self._create_document()
        
        # Build content
        project_info = project_info or {}
        plots = plots or {}
        
        # Cover page
        self.build_cover_page(doc, project_info)
        
        # Executive summary
        self.build_executive_summary(doc, analysis_results)
        
        # Detailed sections
        if 'standards' in analysis_results:
            self.build_standards_section(
                doc,
                analysis_results['standards'],
                plots.get('standards')
            )
        
        if 'blanks' in analysis_results:
            self.build_blanks_section(
                doc,
                analysis_results['blanks'],
                plots.get('blanks')
            )
        
        if 'duplicates' in analysis_results:
            self.build_duplicates_section(
                doc,
                analysis_results['duplicates'],
                plots.get('duplicates')
            )
        
        # JORC section
        self.build_jorc_section(doc, analysis_results)
        
        # Save document
        doc.save(filename)
        
        return filename


# Backwards compatibility alias
DOCXReporter = QAQCDOCXReporter
